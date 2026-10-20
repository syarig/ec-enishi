"""eBay輸出の利益計算・リサーチシートを生成する。

使い方: python3 ebay-export/build_sheet.py ebay-export/ebay-export-sheet.xlsx
（数式の計算結果は Excel で開いたときに計算される）
"""
import datetime
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]

FONT_NAME = "游ゴシック"
BLUE = "0000FF"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")
CALC_FILL = PatternFill("solid", fgColor="F2F2F2")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
NG_FILL = PatternFill("solid", fgColor="FFC7CE")
MID_FILL = PatternFill("solid", fgColor="FFEB9C")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

YEN = '#,##0;[Red]-#,##0;"-"'
USD = '$#,##0.00;[Red]-$#,##0.00;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
INT = '#,##0;[Red]-#,##0;"-"'


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT_NAME, bold=bold, color=color, size=size, italic=italic)


def header_row(ws, row, headers, height=36):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = f(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = height


def style_input(c, fmt=None):
    c.font = f(color=BLUE)
    c.fill = INPUT_FILL
    c.border = BORDER
    if fmt:
        c.number_format = fmt


def style_calc(c, fmt=None, grey=False):
    c.font = f(color="808080" if grey else "000000")
    c.fill = CALC_FILL if grey else PatternFill(fill_type=None)
    c.border = BORDER
    if fmt:
        c.number_format = fmt


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = f(bold=True, size=14)
    if note:
        ws["A2"] = note
        ws["A2"].font = f(color="595959")


wb = Workbook()

# ---------------------------------------------------------------- 設定
st = wb.active
st.title = "設定"
title(st, "設定（為替・手数料・関税・判定基準）",
      "青字・黄色のセルを最新の値に更新してください。灰色は自動計算です。数値は2026年10月時点の目安で、必ず公式情報で確認してください。")
header_row(st, 3, ["項目", "", "値", "単位", "説明・出典"])
st.merge_cells("A3:B3")

params = [
    # row, label, value, fmt, unit, note
    (4, "為替レート（円/ドル）", 150, '#,##0.00', "円", "仮の値。作業する日のレートに更新する。"),
    (5, "為替レート（円/ユーロ）", 165, '#,##0.00', "円", "仮の値。EU向けの固定関税（€3）の円換算に使う。"),
    (6, "eBay 落札手数料率", 0.1325, '0.00%', "%", "多くのカテゴリーの標準料率を仮置き。カテゴリー・ストア有無で変わる。米国では売上税込みの金額にかかるため、実際は少し高くなる。要確認：eBay公式「販売手数料の概要」"),
    (7, "eBay 海外決済手数料率", 0.0165, '0.00%', "%", "International fee。日本セラー（グローバルセラー）に適用。要確認：eBay公式「販売手数料の概要」"),
    (8, "1注文あたり固定手数料", 0.40, '$#,##0.00', "$", "Per-order fee。注文金額により変わる。要確認。"),
    (9, "広告費率（Promoted Listings）", 0, '0.00%', "%", "広告を使う場合だけ入力（例：2〜5%）。使わなければ 0。"),
    (10, "換金コスト率（ドル→円）", 0.02, '0.00%', "%", "売上を円で受け取るときの為替コストの仮置き。受取方法（eBay直接／Payoneer等）で変わる。"),
    (11, "梱包資材費（1件あたり）", 100, '#,##0', "円", "緩衝材・袋・テープなど。段ボールを再利用するなら下げてよい。"),
    (12, "目標利益率", 0.20, '0.0%', "%", "これ以上なら出品候補。計画の基準は 20%。"),
    (13, "最低利益額（1件あたり）", 1000, '#,##0', "円", "これ未満は手間に見合わないので見送り。"),
    (14, "最低販売実績（直近90日）", 3, '#,##0', "件", "Product Research（Terapeak）で確認した販売数の基準。"),
]
for r, label, val, fmt, unit, note in params:
    st.cell(row=r, column=1, value=label).font = f()
    st.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c = st.cell(row=r, column=3, value=val)
    style_input(c, fmt)
    st.cell(row=r, column=4, value=unit).font = f()
    n = st.cell(row=r, column=5, value=note)
    n.font = f(color="595959")
    n.alignment = Alignment(wrap_text=True, vertical="top")

st["A15"] = "（自動）手数料・換金後に残る率"
st["A15"].font = f(color="808080")
st.merge_cells("A15:B15")
st["C15"] = "=(1-($C$6+$C$7+$C$9))*(1-$C$10)"
style_calc(st["C15"], '0.00%', grey=True)
st["E15"] = "最低販売価格の計算に使う。触らない。"
st["E15"].font = f(color="808080")
st["A16"] = "（自動）上の率 − 目標利益率"
st["A16"].font = f(color="808080")
st.merge_cells("A16:B16")
st["C16"] = "=$C$15-$C$12"
style_calc(st["C16"], '0.00%', grey=True)
st["E16"] = "最低販売価格の計算に使う。触らない。"
st["E16"].font = f(color="808080")

# 発送先別の関税
st["A18"] = "発送先別の関税・前払い手数料"
st["A18"].font = f(bold=True, size=11)
header_row(st, 19, ["発送先コード", "国名", "関税率（商品代に対して）", "固定関税（円）", "前払い・通関手数料（円）", "備考"], height=36)
duty_rows = [
    ("US", "米国", 0.15, 0, 0, "日本からの商品にかかる関税率を仮置き（2025年の日米合意の相互関税率）。売り手負担（DDP）前提。最新の率と前払い手数料は eBay「米国関税に関する案内」で要確認。"),
    ("UK", "英国", 0, 0, 0, "£135以下の付加価値税は eBay が購入時に買い手から徴収するので計算に含めない。"),
    ("DE", "ドイツ（EU）", 0, "='設定'!$C$5*3", 0, "2026年7月から€150以下の荷物にも品目ごとに€3の関税。保守的に売り手負担で計算。付加価値税は eBay が買い手から徴収。"),
    ("AU", "オーストラリア", 0, 0, 0, "AUD1,000以下のGSTは eBay が購入時に買い手から徴収するので計算に含めない。"),
]
DUTY_FIRST, DUTY_LAST = 20, 23
for i, (code, name, rate, fixed, fee, note) in enumerate(duty_rows):
    r = DUTY_FIRST + i
    a = st.cell(row=r, column=1, value=code)
    a.font = f(bold=True)
    a.border = BORDER
    a.alignment = Alignment(horizontal="center")
    b = st.cell(row=r, column=2, value=name)
    b.font = f()
    b.border = BORDER
    style_input(st.cell(row=r, column=3, value=rate), '0.0%')
    d = st.cell(row=r, column=4, value=fixed)
    if isinstance(fixed, str):
        style_calc(d, '#,##0')
    else:
        style_input(d, '#,##0')
    style_input(st.cell(row=r, column=5, value=fee), '#,##0')
    n = st.cell(row=r, column=6, value=note)
    n.font = f(color="595959")
    n.alignment = Alignment(wrap_text=True, vertical="top")
    st.row_dimensions[r].height = 42

# 売れた理由タグ
st["A25"] = "売れた理由タグ（プルダウンの選択肢。自由に書き換え・追加してよい）"
st["A25"].font = f(bold=True, size=11)
tags = ["キャラクター人気", "日本限定", "ブランド人気", "海外で入手困難", "サイズ・容量がちょうどいい",
        "セット内容がいい", "ギフト需要", "季節・イベント", "消耗品（リピート）", "その他"]
TAG_FIRST = 26
TAG_LAST = TAG_FIRST + len(tags) - 1
for i, t in enumerate(tags):
    style_input(st.cell(row=TAG_FIRST + i, column=1, value=t))
    st.merge_cells(start_row=TAG_FIRST + i, start_column=1, end_row=TAG_FIRST + i, end_column=2)

st.column_dimensions["A"].width = 18
st.column_dimensions["B"].width = 16
st.column_dimensions["C"].width = 16
st.column_dimensions["D"].width = 14
st.column_dimensions["E"].width = 22
st.column_dimensions["F"].width = 70
for r in range(4, 17):
    st.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
    st.row_dimensions[r].height = 30 if r in (6, 7) else 20
st.row_dimensions[6].height = 44
st.merge_cells("E3:F3")

# ---------------------------------------------------------------- 送料表
sp = wb.create_sheet("送料表")
title(sp, "国際送料表（SpeedPAK Economy 想定）",
      "【重要】下の料金はすべて仮の値です（国別の差も未反映）。eBay 公式の SpeedPAK Economy 料金表で上書きしてください。")
sp["A2"].font = f(bold=True, color="C00000")
header_row(sp, 4, ["重量下限（g）", "重量上限（g）", "US", "UK", "DE", "AU"])
uppers = [100, 200, 300, 400, 500, 600, 700, 800, 900, 1000, 1250, 1500, 1750, 2000]
rates = [1150, 1300, 1450, 1600, 1750, 1900, 2050, 2200, 2350, 2500, 2900, 3300, 3700, 4100]
SP_FIRST = 5
SP_LAST = SP_FIRST + len(uppers) - 1
for i, (u, rate) in enumerate(zip(uppers, rates)):
    r = SP_FIRST + i
    lo = sp.cell(row=r, column=1, value=0 if i == 0 else f"=B{r-1}+1")
    style_calc(lo, '#,##0', grey=True)
    style_input(sp.cell(row=r, column=2, value=u), '#,##0')
    for col in range(3, 7):
        style_input(sp.cell(row=r, column=col, value=rate), '#,##0')
sp.cell(row=SP_LAST + 2, column=1,
        value="・重量は梱包後の重さ（g）。上限を超える重さは「利益計算」で「範囲外」と表示されます。行を増やす場合は利益計算の数式の範囲も広げてください。").font = f(color="595959")
sp.cell(row=SP_LAST + 3, column=1,
        value="・出典：SpeedPAK Economy は「1,141円から」（ネットショップ担当者フォーラム 2024年8月の記事）。それ以外の数値は仮置き。").font = f(color="595959")
for col, w in zip("ABCDEF", [14, 14, 12, 12, 12, 12]):
    sp.column_dimensions[col].width = w
sp.freeze_panes = "A5"

# ---------------------------------------------------------------- 利益計算
pc = wb.create_sheet("利益計算")
title(pc, "利益計算（1行 = 1商品）",
      "青字・黄色のセルを入力すると、手数料・送料・関税を引いた利益と判定が出ます。No.1 の行は例（ツイートの水筒：3,300円 → 約9,000円）。")
pc["B3"] = "◎ 出品候補の件数"
pc["B3"].font = f(bold=True)
pc["B3"].alignment = Alignment(horizontal="right")

headers = [
    "No", "商品名", "仕入れ先", "仕入れ値（円）", "仕入れ時の送料など（円）",
    "販売価格（$）", "買い手負担の送料（$）\n送料無料なら0", "発送先", "梱包後の重量（g）",
    "直近90日の販売数（件）", "売れた理由タグ",
    "売上合計（円）", "eBay手数料（円）", "換金コスト（円）", "国際送料（円）", "関税など（円）",
    "梱包資材費（円）", "総コスト（円）", "利益（円）", "利益率",
    "目標を満たす最低販売価格（$）", "判定", "メモ",
    "（計算用）関税率", "（計算用）固定関税＋手数料（円）",
]
header_row(pc, 4, headers, height=48)
FIRST, LAST = 5, 104
pc["C3"] = f'=COUNTIF(V{FIRST}:V{LAST},"◎*")'
pc["C3"].font = f(bold=True)
pc["C3"].number_format = '#,##0"件"'

S = "'設定'!"
SPS = "'送料表'!"
FX = f"{S}$C$4"
for r in range(FIRST, LAST + 1):
    pc.cell(row=r, column=1, value=r - FIRST + 1).font = f(color="808080")
    pc.cell(row=r, column=1).border = BORDER
    pc.cell(row=r, column=1).alignment = Alignment(horizontal="center")
    for col, fmt in [(2, None), (3, None), (4, YEN), (5, YEN), (6, USD), (7, USD), (8, None),
                     (9, INT), (10, INT), (11, None), (23, None)]:
        style_input(pc.cell(row=r, column=col), fmt)

    ready = f"AND($D{r}<>\"\",$F{r}<>\"\",$H{r}<>\"\",$I{r}<>\"\")"
    formulas = {
        12: (f"=IF({ready},($F{r}+N($G{r}))*{FX},\"\")", YEN),
        13: (f"=IF($L{r}=\"\",\"\",(($F{r}+N($G{r}))*({S}$C$6+{S}$C$7+{S}$C$9)+{S}$C$8)*{FX})", YEN),
        14: (f"=IF($L{r}=\"\",\"\",($L{r}-$M{r})*{S}$C$10)", YEN),
        15: (f"=IF($L{r}=\"\",\"\",IF($I{r}>MAX({SPS}$B${SP_FIRST}:$B${SP_LAST}),\"範囲外\","
             f"INDEX({SPS}$C${SP_FIRST}:$F${SP_LAST},MATCH($I{r},{SPS}$A${SP_FIRST}:$A${SP_LAST},1),"
             f"MATCH($H{r},{SPS}$C$4:$F$4,0))))", YEN),
        16: (f"=IF($L{r}=\"\",\"\",$F{r}*{FX}*$X{r}+$Y{r})", YEN),
        17: (f"=IF($L{r}=\"\",\"\",{S}$C$11)", YEN),
        18: (f"=IF(OR($L{r}=\"\",NOT(ISNUMBER($O{r}))),\"\",N($D{r})+N($E{r})+$M{r}+$N{r}+$O{r}+$P{r}+$Q{r})", YEN),
        19: (f"=IF($R{r}=\"\",\"\",$L{r}-$R{r})", YEN),
        20: (f"=IF(OR($S{r}=\"\",N($L{r})=0),\"\",$S{r}/$L{r})", PCT),
        21: (
            f"=IF(OR($L{r}=\"\",NOT(ISNUMBER($O{r}))),\"\","
            f"IF({S}$C$16-$X{r}<=0,\"設定を確認\","
            f"MAX(({S}$C$8*{FX}*(1-{S}$C$10)+$Y{r}+N($D{r})+N($E{r})+$O{r}+$Q{r}-N($G{r})*{FX}*{S}$C$16)/({FX}*({S}$C$16-$X{r})),"
            f"({S}$C$8*{FX}*(1-{S}$C$10)+$Y{r}+N($D{r})+N($E{r})+$O{r}+$Q{r}+{S}$C$13-N($G{r})*{FX}*{S}$C$15)/({FX}*({S}$C$15-$X{r})))))",
            USD),
        22: (
            f"=IF($L{r}=\"\",\"\",IF(NOT(ISNUMBER($O{r})),\"要確認：重量が送料表の範囲外\","
            f"IF(AND($T{r}>={S}$C$12,$S{r}>={S}$C$13,N($J{r})>={S}$C$14),\"◎ 出品候補\","
            f"IF(AND($T{r}>={S}$C$12,$S{r}>={S}$C$13),\"△ 利益OK・販売実績が不足\",\"× 見送り\"))))",
            None),
        24: (f"=IF($H{r}=\"\",\"\",INDEX({S}$C${DUTY_FIRST}:$C${DUTY_LAST},MATCH($H{r},{S}$A${DUTY_FIRST}:$A${DUTY_LAST},0)))", PCT),
        25: (f"=IF($H{r}=\"\",\"\",INDEX({S}$D${DUTY_FIRST}:$D${DUTY_LAST},MATCH($H{r},{S}$A${DUTY_FIRST}:$A${DUTY_LAST},0))"
             f"+INDEX({S}$E${DUTY_FIRST}:$E${DUTY_LAST},MATCH($H{r},{S}$A${DUTY_FIRST}:$A${DUTY_LAST},0)))", YEN),
    }
    for col, (formula, fmt) in formulas.items():
        c = pc.cell(row=r, column=col, value=formula)
        style_calc(c, fmt, grey=col >= 24)
    pc.cell(row=r, column=19).font = f(bold=True)
    pc.cell(row=r, column=22).font = f(bold=True)

example = {2: "水筒（ツイートの例：3,300円 → 約9,000円）", 3: "Amazon", 4: 3300, 5: 0, 6: 60, 7: 0,
           8: "US", 9: 450, 10: 5, 11: "キャラクター人気",
           23: "例の行。米国の関税と送料で利益が薄くなる。最低販売価格の列で、いくらなら合格かが分かる。"}
for col, v in example.items():
    pc.cell(row=FIRST, column=col, value=v)

dv_dest = DataValidation(type="list", formula1=f"{S}$A${DUTY_FIRST}:$A${DUTY_LAST}", allow_blank=True,
                         showErrorMessage=True, errorTitle="発送先", error="US / UK / DE / AU から選んでください。")
dv_tag = DataValidation(type="list", formula1=f"{S}$A${TAG_FIRST}:$A${TAG_LAST}", allow_blank=True)
pc.add_data_validation(dv_dest)
pc.add_data_validation(dv_tag)
dv_dest.add(f"H{FIRST}:H{LAST}")
dv_tag.add(f"K{FIRST}:K{LAST}")

rng = f"V{FIRST}:V{LAST}"
pc.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($V{FIRST},1)="◎"'], fill=OK_FILL))
pc.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($V{FIRST},1)="△"'], fill=MID_FILL))
pc.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(LEFT($V{FIRST},1)="×",LEFT($V{FIRST},1)="要")'], fill=NG_FILL))

widths = [5, 34, 12, 12, 13, 11, 13, 9, 11, 12, 18, 12, 12, 11, 11, 11, 11, 12, 12, 9, 14, 24, 40, 11, 14]
for i, w in enumerate(widths, start=1):
    pc.column_dimensions[get_column_letter(i)].width = w
pc.freeze_panes = "C5"
pc.auto_filter.ref = f"A4:W{LAST}"

# ---------------------------------------------------------------- ショップ分析
sh = wb.create_sheet("ショップ分析")
title(sh, "① 売れているショップを見る（3〜5店）",
      "Product Research（Terapeak）で「売れたもの（Sold）」「発送元が日本」に絞り、販売数の多い日本人セラーを記録します。")
header_row(sh, 4, ["No", "ショップ名（セラーID）", "URL", "評価数", "主なカテゴリー", "売れ筋上位の商品",
                   "何度も売れている商品", "似た商品を何種類並べているか", "最近増えた商品", "気づき・メモ"], height=40)
SH_FIRST, SH_LAST = 5, 14
for r in range(SH_FIRST, SH_LAST + 1):
    c = sh.cell(row=r, column=1, value=r - SH_FIRST + 1)
    c.font = f(color="808080")
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="top")
    for col in range(2, 11):
        cell = sh.cell(row=r, column=col)
        style_input(cell, INT if col == 4 else None)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sh.row_dimensions[r].height = 45
shop_example = {2: "（例）japan-goods-xxxx", 3: "https://www.ebay.com/str/xxxx", 4: 2500,
                5: "キッチン用品・キャラクター雑貨", 6: "キャラ水筒、弁当箱、箸セット",
                7: "キャラ水筒（色違いで複数回）", 8: "水筒だけで色・容量違い 15種類",
                9: "季節限定デザインの弁当箱", 10: "キャラの日用品を幅広く並べている。容量違いもそれぞれ売れている。"}
for col, v in shop_example.items():
    sh.cell(row=SH_FIRST, column=col, value=v)

sh.cell(row=SH_LAST + 2, column=1, value="共通点のまとめ（3〜5店見たら書く）").font = f(bold=True, size=11)
for i, q in enumerate(["どのカテゴリーが共通して売れているか", "どの価格帯が多いか",
                       "共通する「売れた理由」は何か", "次に候補を広げる元の商品（→「候補展開」へ）"]):
    r = SH_LAST + 3 + i
    q_cell = sh.cell(row=r, column=1, value=q)
    q_cell.font = f()
    q_cell.alignment = Alignment(wrap_text=True, vertical="top")
    sh.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    a = sh.cell(row=r, column=4)
    style_input(a)
    a.alignment = Alignment(wrap_text=True, vertical="top")
    sh.merge_cells(start_row=r, start_column=4, end_row=r, end_column=10)
    for col in range(5, 11):
        sh.cell(row=r, column=col).border = BORDER
    sh.row_dimensions[r].height = 30
for i, w in enumerate([5, 22, 26, 9, 18, 26, 26, 20, 22, 36], start=1):
    sh.column_dimensions[get_column_letter(i)].width = w
sh.freeze_panes = "C5"

# ---------------------------------------------------------------- 候補展開
cx = wb.create_sheet("候補展開")
title(cx, "② 横・縦・周辺に広げる → ③ 売れた理由で絞る",
      "1つの売れ筋から20〜30件の候補を出し、元の商品と「売れた理由」が同じものだけ残します。残ったものを「利益計算」に転記します。")
header_row(cx, 4, ["No", "元の商品（売れている商品）", "元の商品が売れた理由", "広げ方", "広げた種類",
                   "候補の商品名", "候補も同じ理由を満たす？", "直近90日の販売数（件）", "利益計算に転記", "メモ"], height=40)
CX_FIRST, CX_LAST = 5, 204
for r in range(CX_FIRST, CX_LAST + 1):
    c = cx.cell(row=r, column=1, value=r - CX_FIRST + 1)
    c.font = f(color="808080")
    c.border = BORDER
    c.alignment = Alignment(horizontal="center")
    for col in range(2, 11):
        style_input(cx.cell(row=r, column=col), INT if col == 8 else None)
cx_examples = [
    ["キャラクター水筒 480ml", "キャラクター人気", "横", "容量違い", "同じキャラの水筒 800ml", "○", 4, "○", ""],
    ["キャラクター水筒 480ml", "キャラクター人気", "縦", "同シリーズ", "同じキャラの弁当箱", "○", 6, "○", ""],
    ["キャラクター水筒 480ml", "キャラクター人気", "周辺", "交換品・消耗品", "この水筒の替えパッキン（キャラなし）", "×", 2, "", "キャラ人気が理由なので、キャラなしの部品は優先度を下げる"],
]
for i, row in enumerate(cx_examples):
    for j, v in enumerate(row, start=2):
        cx.cell(row=CX_FIRST + i, column=j, value=v)

dv_dir = DataValidation(type="list", formula1='"横,縦,周辺"', allow_blank=True)
dv_reason = DataValidation(type="list", formula1=f"{S}$A${TAG_FIRST}:$A${TAG_LAST}", allow_blank=True)
dv_match = DataValidation(type="list", formula1='"○,×,？"', allow_blank=True)
dv_done = DataValidation(type="list", formula1='"○"', allow_blank=True)
for dv, col in [(dv_dir, "D"), (dv_reason, "C"), (dv_match, "G"), (dv_done, "I")]:
    cx.add_data_validation(dv)
    dv.add(f"{col}{CX_FIRST}:{col}{CX_LAST}")
cx.conditional_formatting.add(f"G{CX_FIRST}:G{CX_LAST}", FormulaRule(formula=[f'$G{CX_FIRST}="○"'], fill=OK_FILL))
cx.conditional_formatting.add(f"G{CX_FIRST}:G{CX_LAST}", FormulaRule(formula=[f'$G{CX_FIRST}="×"'], fill=NG_FILL))

# 広げ方の早見表（ツイートの②）
ref_col = 12
cx.cell(row=4, column=ref_col, value="広げ方の早見表").font = f(bold=True, size=11)
ref = {
    "横（同じ商品の別バリエーション）": ["色違い", "サイズ違い", "容量違い", "個数違い", "セット違い", "デザイン違い", "キャラクター違い", "季節違い"],
    "縦（同ブランド・同シリーズ）": ["同じブランドの別商品", "上位モデル", "下位モデル", "旧モデル", "新モデル", "同シリーズ", "同じ価格帯", "高価格帯・低価格帯"],
    "周辺（一緒に使うもの）": ["一緒に使う商品", "交換品", "消耗品", "詰め替え", "ケース", "カバー", "アクセサリー", "収納用品", "メンテナンス用品", "ギフト需要の商品"],
}
for j, (head, items) in enumerate(ref.items()):
    hc = cx.cell(row=5, column=ref_col + j, value=head)
    hc.font = f(bold=True)
    hc.fill = SUB_FILL
    hc.border = BORDER
    hc.alignment = Alignment(wrap_text=True)
    for i, it in enumerate(items):
        ic = cx.cell(row=6 + i, column=ref_col + j, value=it)
        ic.font = f()
        ic.border = BORDER
cx.row_dimensions[5].height = 30
for i, w in enumerate([5, 24, 20, 8, 16, 30, 12, 12, 10, 34, 3, 22, 22, 22], start=1):
    cx.column_dimensions[get_column_letter(i)].width = w
cx.freeze_panes = "C5"

# ---------------------------------------------------------------- 実績記録
lg = wb.create_sheet("実績記録")
title(lg, "実績記録（毎日つける → 月ごとの時給が出る）",
      "作業時間も必ず記録します。右の表に月ごとの合計と時給が自動で出ます。")
header_row(lg, 4, ["日付", "作業時間（h）", "出品数", "販売数", "確定利益（円）", "メモ"])
LG_FIRST, LG_LAST = 5, 404
for r in range(LG_FIRST, LG_LAST + 1):
    style_input(lg.cell(row=r, column=1), "yyyy/mm/dd")
    style_input(lg.cell(row=r, column=2), '0.0')
    style_input(lg.cell(row=r, column=3), INT)
    style_input(lg.cell(row=r, column=4), INT)
    style_input(lg.cell(row=r, column=5), YEN)
    style_input(lg.cell(row=r, column=6))
for col, v in enumerate([datetime.date(2026, 10, 20), 2.5, 0, 0, 0, "（例）eBayアカウント作成・古物商許可の書類準備"], start=1):
    lg.cell(row=LG_FIRST, column=col, value=v)

sum_col = 8  # H
sum_headers = ["月", "作業時間（h）", "出品数", "販売数", "確定利益（円）", "時給（円/h）", "計画の目安"]
for i, h in enumerate(sum_headers):
    c = lg.cell(row=4, column=sum_col + i, value=h)
    c.font = f(bold=True, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
plan = {0: "準備・不用品で最初の売上・評価5〜10件", 2: "月10〜20件販売・月の利益1〜2万円",
        5: "出品150〜200品・月30件・月の利益4.5万円"}
MONTHS = 7
for i in range(MONTHS):
    r = 5 + i
    m = lg.cell(row=r, column=sum_col, value=datetime.date(2026, 10, 1) if i == 0 else f"=DATE(YEAR(H{r-1}),MONTH(H{r-1})+1,1)")
    if i == 0:
        style_input(m, 'yyyy"年"m"月"')
    else:
        style_calc(m, 'yyyy"年"m"月"')
    rng_a = f"$A${LG_FIRST}:$A${LG_LAST}"
    cond = f'{rng_a},">="&$H{r},{rng_a},"<"&DATE(YEAR($H{r}),MONTH($H{r})+1,1)'
    for j, src in enumerate(["B", "C", "D", "E"]):
        fmt = '0.0' if src == "B" else (YEN if src == "E" else INT)
        c = lg.cell(row=r, column=sum_col + 1 + j, value=f"=SUMIFS(${src}${LG_FIRST}:${src}${LG_LAST},{cond})")
        style_calc(c, fmt)
    hw = lg.cell(row=r, column=sum_col + 5, value=f"=IF(I{r}=0,\"\",L{r}/I{r})")
    style_calc(hw, YEN)
    hw.font = f(bold=True)
    p = lg.cell(row=r, column=sum_col + 6, value=(f"{i+1}ヶ月目：" + plan[i]) if i in plan else f"{i+1}ヶ月目")
    p.font = f(color="595959")
    p.border = BORDER
tr = 5 + MONTHS
lg.cell(row=tr, column=sum_col, value="合計").font = f(bold=True)
lg.cell(row=tr, column=sum_col).border = BORDER
for j, col in enumerate(["I", "J", "K", "L"]):
    fmt = '0.0' if col == "I" else (YEN if col == "L" else INT)
    c = lg.cell(row=tr, column=sum_col + 1 + j, value=f"=SUM({col}5:{col}{tr-1})")
    style_calc(c, fmt)
    c.font = f(bold=True)
c = lg.cell(row=tr, column=sum_col + 5, value=f"=IF(I{tr}=0,\"\",L{tr}/I{tr})")
style_calc(c, YEN)
c.font = f(bold=True)
lg.cell(row=tr + 1, column=sum_col, value="※ 1ヶ月目の月は H5 を変えると、以降の月が自動でずれます。").font = f(color="595959")
for col, w in zip("ABCDEFGHIJKLMN", [12, 11, 9, 9, 13, 36, 3, 12, 11, 9, 9, 13, 12, 40]):
    lg.column_dimensions[col].width = w
lg.freeze_panes = "A5"

# ---------------------------------------------------------------- 使い方
hw = wb.create_sheet("使い方", 0)
hw.column_dimensions["A"].width = 4
hw.column_dimensions["B"].width = 110
hw["A1"] = "eBay輸出 利益計算・リサーチシート"
hw["A1"].font = f(bold=True, size=16)
lines = [
    ("h", "セルの色"),
    ("legend_in", "青字・黄色背景 … 入力するセル"),
    ("legend_calc", "黒字・白背景 … 自動計算（触らない）"),
    ("legend_grey", "灰色背景 … 計算用の補助セル（触らない）"),
    ("", ""),
    ("h", "使う順番"),
    ("", "1. 「設定」… 為替レートを今日の値に更新。手数料率・関税率は eBay 公式で確認して直す（最初に1回、あとは月1回）"),
    ("", "2. 「送料表」… 仮の値が入っています。SpeedPAK Economy の公式料金で必ず上書きする"),
    ("", "3. 「ショップ分析」… ① 売れている日本人セラーを3〜5店記録し、共通点をまとめる"),
    ("", "4. 「候補展開」… ② 1つの売れ筋を横・縦・周辺に20〜30件広げ、③ 売れた理由が同じものに○を付ける"),
    ("", "5. 「利益計算」… ○の候補を転記。判定が「◎ 出品候補」のものだけ出品する"),
    ("", "    「目標を満たす最低販売価格」の列で、いくらで売れれば合格かが分かる（相場がそれより下なら見送り）"),
    ("", "6. 「実績記録」… 毎日の作業時間・出品数・販売数・利益を記録。月ごとの時給を見る"),
    ("", ""),
    ("h", "判定の基準（「設定」で変更できる）"),
    ("", "◎ 出品候補 … 利益率20%以上 かつ 利益1,000円以上 かつ 直近90日に3件以上売れている"),
    ("", "△ 利益OK・販売実績が不足 … 利益は出るが売れた実績が少ない。様子見"),
    ("", "× 見送り … 利益率か利益額が足りない"),
    ("", ""),
    ("h", "計算の前提と注意"),
    ("", "・利益 = 売上 −（仕入れ値 + 仕入れ時の送料 + eBay手数料 + 換金コスト + 国際送料 + 関税など + 梱包資材費）"),
    ("", "・eBay手数料 =（販売価格 + 買い手負担の送料）×（落札手数料率 + 海外決済手数料率 + 広告費率）+ 1注文あたり固定手数料"),
    ("", "・関税は売り手負担（DDP）前提の保守的な計算。英国・EU・オーストラリアの付加価値税（GST）は eBay が買い手から徴収するので含めていない"),
    ("", "・米国では売上税込みの金額に落札手数料がかかるため、実際の手数料は計算より少し高くなる"),
    ("", "・手数料率・関税率・送料は変わります。数値は2026年10月時点の目安なので、作業前に公式情報で確認してください"),
    ("", ""),
    ("h", "出典"),
    ("", "eBay 販売手数料の概要：https://wwwcojp.ebay.com/fee/overview-selling-fees"),
    ("", "eBay 米国関税に関する案内：https://wwwcojp2.ebay.com/tariffs"),
    ("", "eBay DDP（関税元払い）案内：https://wwwcojp.ebay.com/ddp-info"),
    ("", "SpeedPAK Economy の概要（ネットショップ担当者フォーラム）：https://netshop.impress.co.jp/node/12664"),
    ("", "日本郵便 米国宛て郵便物の引受再開（2026年4月）：https://www.post.japanpost.jp/int/information/2026/0413_01.html"),
]
r = 3
for kind, text in lines:
    if kind == "h":
        c = hw.cell(row=r, column=1, value=text)
        c.font = f(bold=True, size=12)
    elif kind.startswith("legend"):
        c = hw.cell(row=r, column=2, value=text)
        if kind == "legend_in":
            c.font = f(color=BLUE)
            c.fill = INPUT_FILL
        elif kind == "legend_grey":
            c.font = f(color="808080")
            c.fill = CALC_FILL
        else:
            c.font = f()
        c.border = BORDER
    else:
        c = hw.cell(row=r, column=2, value=text)
        c.font = f()
        c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1

for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

wb.active = 0
wb.save(OUT)
print("saved", OUT)
